from pathlib import Path

import pytest
from docx import Document
from openpyxl import Workbook, load_workbook

from agent.plugins import excel_processor, web_query, word_processor


def _create_excel_file(path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Sheet1"
    sheet.append(["部门", "金额", "状态"])
    sheet.append(["销售", 100, "已完成"])
    sheet.append(["采购", 50, "处理中"])
    sheet.append(["销售", 80, "已完成"])
    workbook.save(path)


def _create_word_file(path: Path) -> None:
    document = Document()
    document.add_heading("项目背景介绍", level=1)
    document.add_paragraph("项目目标是提升流程效率。")
    document.add_paragraph("关键字需要被替换。")
    document.add_paragraph("模板变量：{{company}}")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "字段"
    table.cell(0, 1).text = "值"
    table.cell(1, 0).text = "负责人"
    table.cell(1, 1).text = "Alice"
    document.save(path)


@pytest.mark.asyncio
async def test_excel_processor_returns_real_summary(tmp_path):
    file_path = tmp_path / "sales.xlsx"
    _create_excel_file(file_path)

    result = await excel_processor.run("summary", file_path=str(file_path), sheet_name="Sheet1")

    assert result["status"] == "success"
    assert result["data"]["row_count"] == 3
    assert result["data"]["numeric_summary"]["金额"]["sum"] == 230.0


@pytest.mark.asyncio
async def test_excel_processor_creates_report_file(tmp_path):
    file_path = tmp_path / "sales.xlsx"
    report_path = tmp_path / "sales_report.xlsx"
    _create_excel_file(file_path)

    result = await excel_processor.run("create_report", file_path=str(file_path), output_path=str(report_path))

    assert result["status"] == "success"
    assert report_path.exists()
    workbook = load_workbook(report_path)
    assert workbook.active["A2"].value == "Source File"


@pytest.mark.asyncio
async def test_excel_processor_reads_range_preview(tmp_path):
    file_path = tmp_path / "sales.xlsx"
    _create_excel_file(file_path)

    result = await excel_processor.run("read_range", file_path=str(file_path), range_ref="A1:B3")

    assert result["status"] == "success"
    assert result["data"]["range_ref"] == "A1:B3"
    assert result["data"]["preview"][0]["部门"] == "销售"


@pytest.mark.asyncio
async def test_excel_processor_writes_cells_to_output_file(tmp_path):
    file_path = tmp_path / "sales.xlsx"
    output_path = tmp_path / "sales_updated.xlsx"
    _create_excel_file(file_path)

    result = await excel_processor.run(
        "write_cells",
        file_path=str(file_path),
        output_path=str(output_path),
        cell_updates=[{"cell": "B2", "value": 200}],
    )

    assert result["status"] == "success"
    workbook = load_workbook(output_path)
    assert workbook.active["B2"].value == 200


@pytest.mark.asyncio
async def test_excel_processor_appends_rows_to_output_file(tmp_path):
    file_path = tmp_path / "sales.xlsx"
    output_path = tmp_path / "sales_appended.xlsx"
    _create_excel_file(file_path)

    result = await excel_processor.run(
        "append_rows",
        file_path=str(file_path),
        output_path=str(output_path),
        rows=[{"部门": "财务", "金额": 120, "状态": "已完成"}],
    )

    assert result["status"] == "success"
    workbook = load_workbook(output_path)
    assert workbook.active.max_row == 5
    assert workbook.active["A5"].value == "财务"


@pytest.mark.asyncio
async def test_excel_processor_aggregates_rows_by_group(tmp_path):
    file_path = tmp_path / "sales.xlsx"
    _create_excel_file(file_path)

    result = await excel_processor.run(
        "aggregate",
        file_path=str(file_path),
        group_by="部门",
        metric_column="金额",
        metric="sum",
    )

    assert result["status"] == "success"
    groups = {item["group"]: item["value"] for item in result["data"]["groups"]}
    assert groups["销售"] == 180.0
    assert groups["采购"] == 50.0


@pytest.mark.asyncio
async def test_word_processor_replaces_text_and_saves_output(tmp_path):
    document_path = tmp_path / "demo.docx"
    output_path = tmp_path / "demo_replaced.docx"
    _create_word_file(document_path)

    result = await word_processor.run(
        "replace_text",
        document_path=str(document_path),
        keyword="关键字",
        replacement="术语",
        output_path=str(output_path),
    )

    assert result["status"] == "success"
    assert output_path.exists()
    document = Document(output_path)
    assert any("术语需要被替换" in paragraph.text for paragraph in document.paragraphs)


@pytest.mark.asyncio
async def test_word_processor_summarizes_real_document(tmp_path):
    document_path = tmp_path / "demo.docx"
    _create_word_file(document_path)

    result = await word_processor.run("summarize", document_path=str(document_path))

    assert result["status"] == "success"
    assert result["data"]["paragraph_count"] == 4
    assert "项目背景介绍" in result["data"]["summary"]


@pytest.mark.asyncio
async def test_word_processor_fills_template_and_saves_output(tmp_path):
    document_path = tmp_path / "demo.docx"
    output_path = tmp_path / "demo_filled.docx"
    _create_word_file(document_path)

    result = await word_processor.run(
        "fill_template",
        document_path=str(document_path),
        replacements={"company": "FlowMind"},
        output_path=str(output_path),
    )

    assert result["status"] == "success"
    document = Document(output_path)
    assert any("模板变量：FlowMind" in paragraph.text for paragraph in document.paragraphs)


@pytest.mark.asyncio
async def test_word_processor_extracts_outline(tmp_path):
    document_path = tmp_path / "demo.docx"
    _create_word_file(document_path)

    result = await word_processor.run("extract_outline", document_path=str(document_path))

    assert result["status"] == "success"
    assert result["data"]["count"] == 1
    assert result["data"]["items"][0]["text"] == "项目背景介绍"


@pytest.mark.asyncio
async def test_word_processor_extracts_tables(tmp_path):
    document_path = tmp_path / "demo.docx"
    _create_word_file(document_path)

    result = await word_processor.run("extract_tables", document_path=str(document_path))

    assert result["status"] == "success"
    assert result["data"]["table_count"] == 1
    assert result["data"]["tables"][0]["rows"][1][1] == "Alice"


@pytest.mark.asyncio
async def test_web_query_search_returns_parsed_results(monkeypatch):
    html = """
    <html>
      <body>
        <div class="result">
          <a class="result__a" href="https://example.com/a">Result A</a>
          <a class="result__snippet">Snippet A</a>
        </div>
        <div class="result">
          <a class="result__a" href="https://example.com/b">Result B</a>
          <a class="result__snippet">Snippet B</a>
        </div>
      </body>
    </html>
    """

    class DummyResponse:
        def __init__(self, text):
            self.text = text

        def raise_for_status(self):
            return None

    monkeypatch.setattr("agent.plugins.web_query.requests.get", lambda *args, **kwargs: DummyResponse(html))

    result = await web_query.run(query="天气预报", site="example.com", top_k=2)

    assert result["status"] == "success"
    assert len(result["data"]["results"]) == 2
    assert result["data"]["results"][0]["title"] == "Result A"


@pytest.mark.asyncio
async def test_web_query_fetch_page_returns_page_summary(monkeypatch):
    html = """
    <html>
      <head><title>Demo Page</title></head>
      <body>
        <p>第一段内容。</p>
        <p>第二段内容。</p>
      </body>
    </html>
    """

    class DummyResponse:
        def __init__(self, text):
            self.text = text

        def raise_for_status(self):
            return None

    monkeypatch.setattr("agent.plugins.web_query.requests.get", lambda *args, **kwargs: DummyResponse(html))

    result = await web_query.run(mode="fetch_page", url="https://example.com/article")

    assert result["status"] == "success"
    assert result["data"]["title"] == "Demo Page"
    assert result["data"]["paragraph_count"] == 2


@pytest.mark.asyncio
async def test_web_query_extract_structured_returns_headings_and_links(monkeypatch):
    html = """
    <html>
      <head><title>Demo Page</title></head>
      <body>
        <h1>主标题</h1>
        <a href="https://example.com/a">Link A</a>
        <p>第一段内容。</p>
      </body>
    </html>
    """

    class DummyResponse:
        def __init__(self, text):
            self.text = text

        def raise_for_status(self):
            return None

    monkeypatch.setattr("agent.plugins.web_query.requests.get", lambda *args, **kwargs: DummyResponse(html))

    result = await web_query.run(mode="extract_structured", url="https://example.com/article")

    assert result["status"] == "success"
    assert result["data"]["headings"] == ["主标题"]
    assert result["data"]["links"][0]["href"] == "https://example.com/a"


@pytest.mark.asyncio
async def test_web_query_fetch_many_returns_multiple_pages(monkeypatch):
    pages = {
        "https://example.com/a": "<html><head><title>A</title></head><body><p>Alpha</p></body></html>",
        "https://example.com/b": "<html><head><title>B</title></head><body><p>Beta</p></body></html>",
    }

    class DummyResponse:
        def __init__(self, text):
            self.text = text

        def raise_for_status(self):
            return None

    monkeypatch.setattr(
        "agent.plugins.web_query.requests.get",
        lambda url, *args, **kwargs: DummyResponse(pages[url]),
    )

    result = await web_query.run(mode="fetch_many", urls=list(pages), top_k=2)

    assert result["status"] == "success"
    assert result["data"]["count"] == 2
    assert result["data"]["pages"][0]["title"] == "A"


@pytest.mark.asyncio
async def test_web_query_selector_extract_returns_selected_items(monkeypatch):
    html = """
    <html>
      <body>
        <ul>
          <li class="item">One</li>
          <li class="item">Two</li>
        </ul>
      </body>
    </html>
    """

    class DummyResponse:
        def __init__(self, text):
            self.text = text

        def raise_for_status(self):
            return None

    monkeypatch.setattr("agent.plugins.web_query.requests.get", lambda *args, **kwargs: DummyResponse(html))

    result = await web_query.run(mode="selector_extract", url="https://example.com", selector=".item", limit=2)

    assert result["status"] == "success"
    assert result["data"]["items"] == ["One", "Two"]
