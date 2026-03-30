import asyncio
from pathlib import Path
from statistics import mean
from typing import Any, Dict, Iterable, List

from openpyxl import Workbook, load_workbook

from utils.plugin_result import error_result, success_result

SUPPORTED_OPERATIONS = {
    "summary",
    "filter",
    "create_report",
    "read_range",
    "write_cells",
    "append_rows",
    "aggregate",
}
SUPPORTED_METRICS = {"sum", "avg", "min", "max", "count"}


def _normalize_headers(rows: List[List[Any]]) -> List[str]:
    return [str(value).strip() if value is not None else "" for value in rows[0]]


def _open_sheet(file_path: str, sheet_name: str = "", *, data_only: bool = False):
    path = Path(file_path)
    if not path.exists():
        raise FileNotFoundError(f"Excel file not found: {file_path}")

    workbook = load_workbook(path, data_only=data_only)
    sheet = workbook[sheet_name] if sheet_name and sheet_name in workbook.sheetnames else workbook.active
    return workbook, sheet


def _load_sheet_rows(file_path: str, sheet_name: str = "", range_ref: str = "") -> Dict[str, Any]:
    workbook, sheet = _open_sheet(file_path, sheet_name, data_only=True)
    iterator = sheet[range_ref] if range_ref else sheet.iter_rows()
    rows = [[cell.value for cell in row] for row in iterator]
    rows = [row for row in rows if any(value is not None and str(value).strip() != "" for value in row)]
    workbook.close()
    if not rows:
        raise ValueError("Excel sheet is empty")

    headers = _normalize_headers(rows)
    records = []
    for row in rows[1:]:
        padded = row + [None] * max(0, len(headers) - len(row))
        records.append({headers[index] or f"column_{index + 1}": padded[index] for index in range(len(headers))})

    return {
        "sheet_name": sheet.title,
        "headers": headers,
        "records": records,
    }


def _summarize_numeric_columns(records: List[Dict[str, Any]], headers: List[str]) -> Dict[str, Dict[str, float]]:
    numeric_summary: Dict[str, Dict[str, float]] = {}
    for header in headers:
        values = [record.get(header) for record in records]
        numeric_values = [float(value) for value in values if isinstance(value, (int, float))]
        if numeric_values:
            numeric_summary[header] = {
                "count": len(numeric_values),
                "sum": round(sum(numeric_values), 2),
                "avg": round(mean(numeric_values), 2),
                "min": round(min(numeric_values), 2),
                "max": round(max(numeric_values), 2),
            }
    return numeric_summary


def _write_report(output_path: str, summary: Dict[str, Any]) -> str:
    path = Path(output_path)
    path.parent.mkdir(parents=True, exist_ok=True)

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Summary"
    sheet.append(["Item", "Value"])
    sheet.append(["Source File", summary["source_file"]])
    sheet.append(["Sheet Name", summary["sheet_name"]])
    sheet.append(["Row Count", summary["row_count"]])
    sheet.append(["Column Count", summary["column_count"]])
    sheet.append([])
    sheet.append(["Numeric Column", "count", "sum", "avg", "min", "max"])
    for column, stats in summary["numeric_summary"].items():
        sheet.append([column, stats["count"], stats["sum"], stats["avg"], stats["min"], stats["max"]])

    workbook.save(path)
    workbook.close()
    return str(path)


def _default_output_path(file_path: str, suffix: str) -> str:
    source = Path(file_path)
    return str(source.with_name(f"{source.stem}_{suffix}{source.suffix or '.xlsx'}"))


def _save_workbook(workbook, file_path: str, output_path: str, suffix: str) -> str:
    save_path = output_path or _default_output_path(file_path, suffix)
    path = Path(save_path)
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)
    workbook.close()
    return str(path)


def _read_range(file_path: str, sheet_name: str, range_ref: str) -> Dict[str, Any]:
    dataset = _load_sheet_rows(file_path, sheet_name, range_ref)
    return {
        "operation": "read_range",
        "source_file": Path(file_path).name,
        "sheet_name": dataset["sheet_name"],
        "range_ref": range_ref or "",
        "headers": dataset["headers"],
        "row_count": len(dataset["records"]),
        "preview": dataset["records"][:20],
    }


def _write_cells(file_path: str, sheet_name: str, cell_updates: Iterable[dict], output_path: str) -> Dict[str, Any]:
    workbook, sheet = _open_sheet(file_path, sheet_name)
    updates = list(cell_updates or [])
    for item in updates:
        cell_ref = str((item or {}).get("cell") or "").strip()
        if not cell_ref:
            raise ValueError("Each cell update must include a cell")
        sheet[cell_ref] = (item or {}).get("value")

    saved_path = _save_workbook(workbook, file_path, output_path, "updated")
    return {
        "operation": "write_cells",
        "sheet_name": sheet.title,
        "updated_count": len(updates),
        "output_path": saved_path,
    }


def _coerce_row_values(row_item: Any, headers: List[str]) -> List[Any]:
    if isinstance(row_item, dict):
        return [row_item.get(header) for header in headers]
    if isinstance(row_item, (list, tuple)):
        values = list(row_item)
        return values + [None] * max(0, len(headers) - len(values))
    raise ValueError("rows must contain objects or arrays")


def _append_rows(file_path: str, sheet_name: str, rows: Iterable[Any], output_path: str) -> Dict[str, Any]:
    workbook, sheet = _open_sheet(file_path, sheet_name)
    values = list(rows or [])
    if not values:
        raise ValueError("rows is required for append_rows")

    header_cells = [cell.value for cell in sheet[1]]
    headers = [str(value).strip() if value is not None else f"column_{index + 1}" for index, value in enumerate(header_cells)]
    if not any(headers):
        raise ValueError("Excel sheet is missing a header row")

    extra_headers: List[str] = []
    for row_item in values:
        if isinstance(row_item, dict):
            for key in row_item.keys():
                key_text = str(key).strip()
                if key_text and key_text not in headers and key_text not in extra_headers:
                    extra_headers.append(key_text)

    for header in extra_headers:
        headers.append(header)
        sheet.cell(row=1, column=len(headers), value=header)

    for row_item in values:
        sheet.append(_coerce_row_values(row_item, headers))

    saved_path = _save_workbook(workbook, file_path, output_path, "appended")
    return {
        "operation": "append_rows",
        "sheet_name": sheet.title,
        "appended_count": len(values),
        "headers": headers,
        "output_path": saved_path,
    }


def _aggregate_records(records: List[Dict[str, Any]], group_by: str, metric_column: str, metric: str) -> List[Dict[str, Any]]:
    groups: Dict[str, List[Any]] = {}
    for record in records:
        key = str(record.get(group_by, ""))
        groups.setdefault(key, []).append(record.get(metric_column))

    aggregated: List[Dict[str, Any]] = []
    for key, values in groups.items():
        numeric_values = [float(value) for value in values if isinstance(value, (int, float))]
        if metric == "count":
            metric_value = len(values)
        else:
            if not numeric_values:
                continue
            if metric == "sum":
                metric_value = round(sum(numeric_values), 2)
            elif metric == "avg":
                metric_value = round(mean(numeric_values), 2)
            elif metric == "min":
                metric_value = round(min(numeric_values), 2)
            else:
                metric_value = round(max(numeric_values), 2)
        aggregated.append(
            {
                "group": key,
                "metric": metric,
                "metric_column": metric_column,
                "value": metric_value,
            }
        )
    return aggregated


async def run(
    operation: str,
    file_path: str,
    sheet_name: str = "",
    range_ref: str = "",
    filter_column: str = "",
    filter_value: str = "",
    output_path: str = "",
    cell_updates: List[dict] | None = None,
    rows: List[Any] | None = None,
    group_by: str = "",
    metric_column: str = "",
    metric: str = "sum",
) -> dict:
    if not operation:
        return error_result("operation is required")
    if operation not in SUPPORTED_OPERATIONS:
        return error_result(
            "Unsupported excel operation",
            error=f"operation must be one of {sorted(SUPPORTED_OPERATIONS)}",
        )
    if not file_path:
        return error_result("file_path is required")

    if operation == "write_cells":
        if not cell_updates:
            return error_result("cell_updates is required for write_cells")
        try:
            result = await asyncio.to_thread(_write_cells, file_path, sheet_name, cell_updates, output_path)
            return success_result(data=result)
        except FileNotFoundError as exc:
            return error_result("Excel file does not exist", error=str(exc))
        except Exception as exc:
            return error_result("Failed to write Excel cells", error=str(exc))

    if operation == "append_rows":
        if not rows:
            return error_result("rows is required for append_rows")
        try:
            result = await asyncio.to_thread(_append_rows, file_path, sheet_name, rows, output_path)
            return success_result(data=result)
        except FileNotFoundError as exc:
            return error_result("Excel file does not exist", error=str(exc))
        except Exception as exc:
            return error_result("Failed to append Excel rows", error=str(exc))

    try:
        dataset = await asyncio.to_thread(_load_sheet_rows, file_path, sheet_name, range_ref)
    except FileNotFoundError as exc:
        return error_result("Excel file does not exist", error=str(exc))
    except Exception as exc:
        return error_result("Failed to read Excel file", error=str(exc))

    headers = dataset["headers"]
    records = dataset["records"]
    summary = {
        "operation": operation,
        "source_file": Path(file_path).name,
        "sheet_name": dataset["sheet_name"],
        "headers": headers,
        "row_count": len(records),
        "column_count": len(headers),
        "numeric_summary": _summarize_numeric_columns(records, headers),
    }

    if operation == "summary":
        summary["preview"] = records[:10]
        return success_result(data=summary)

    if operation == "read_range":
        summary["range_ref"] = range_ref or ""
        summary["preview"] = records[:20]
        return success_result(data=summary)

    if operation == "filter":
        if not filter_column or filter_value == "":
            return error_result("filter_column and filter_value are required for filter")
        if filter_column not in headers:
            return error_result("filter_column not found", error=f"Unknown column: {filter_column}")
        filtered = [record for record in records if str(record.get(filter_column, "")) == str(filter_value)]
        summary["filter_column"] = filter_column
        summary["filter_value"] = filter_value
        summary["matched_count"] = len(filtered)
        summary["preview"] = filtered[:20]
        return success_result(data=summary)

    if operation == "aggregate":
        if not group_by or not metric_column:
            return error_result("group_by and metric_column are required for aggregate")
        if group_by not in headers:
            return error_result("group_by not found", error=f"Unknown column: {group_by}")
        if metric_column not in headers:
            return error_result("metric_column not found", error=f"Unknown column: {metric_column}")
        metric_name = str(metric or "sum").strip().lower()
        if metric_name not in SUPPORTED_METRICS:
            return error_result("metric is invalid", error=f"metric must be one of {sorted(SUPPORTED_METRICS)}")
        summary["group_by"] = group_by
        summary["metric_column"] = metric_column
        summary["metric"] = metric_name
        summary["groups"] = _aggregate_records(records, group_by, metric_column, metric_name)
        return success_result(data=summary)

    report_path = output_path or _default_output_path(file_path, "summary")
    try:
        saved_path = await asyncio.to_thread(_write_report, report_path, summary)
    except Exception as exc:
        return error_result("Failed to create Excel report", error=str(exc))

    summary["output_path"] = saved_path
    return success_result(data=summary)
