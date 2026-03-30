import asyncio
import re
from pathlib import Path
from typing import Dict, List

from docx import Document
from docx.document import Document as DocumentType
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

from utils.plugin_result import error_result, success_result

SUPPORTED_OPERATIONS = {
    "summarize",
    "format",
    "replace_text",
    "fill_template",
    "extract_outline",
    "extract_tables",
}


def _load_document(document_path: str) -> DocumentType:
    path = Path(document_path)
    if not path.exists():
        raise FileNotFoundError(f"Word file not found: {document_path}")
    return Document(path)


def _collect_paragraph_text(document: DocumentType) -> List[str]:
    return [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]


def _default_output_path(document_path: str, suffix: str) -> str:
    path = Path(document_path)
    return str(path.with_name(f"{path.stem}_{suffix}.docx"))


def _iter_text_blocks(document: DocumentType):
    for paragraph in document.paragraphs:
        yield paragraph
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                yield cell


def _summarize_document(document_path: str) -> dict:
    document = _load_document(document_path)
    paragraphs = _collect_paragraph_text(document)
    summary_text = " ".join(paragraphs[:3])[:300]
    return {
        "operation": "summarize",
        "document_name": Path(document_path).name,
        "paragraph_count": len(paragraphs),
        "word_count": sum(len(text.split()) for text in paragraphs),
        "summary": summary_text,
        "preview": paragraphs[:5],
    }


def _replace_text(document_path: str, keyword: str, replacement: str, output_path: str) -> dict:
    document = _load_document(document_path)
    replacements = 0

    for paragraph in document.paragraphs:
        if keyword in paragraph.text:
            paragraph.text = paragraph.text.replace(keyword, replacement)
            replacements += 1

    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                if keyword in cell.text:
                    cell.text = cell.text.replace(keyword, replacement)
                    replacements += 1

    save_path = output_path or _default_output_path(document_path, "replaced")
    Path(save_path).parent.mkdir(parents=True, exist_ok=True)
    document.save(save_path)
    return {
        "operation": "replace_text",
        "document_name": Path(document_path).name,
        "keyword": keyword,
        "replacement": replacement,
        "replacements": replacements,
        "output_path": save_path,
    }


def _format_document(document_path: str, output_path: str) -> dict:
    document = _load_document(document_path)
    formatted = 0

    for paragraph in document.paragraphs:
        if not paragraph.text.strip():
            continue
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        for run in paragraph.runs:
            run.font.name = "Calibri"
            run.font.size = Pt(11)
        formatted += 1

    save_path = output_path or _default_output_path(document_path, "formatted")
    Path(save_path).parent.mkdir(parents=True, exist_ok=True)
    document.save(save_path)
    return {
        "operation": "format",
        "document_name": Path(document_path).name,
        "formatted_paragraphs": formatted,
        "output_path": save_path,
    }


def _fill_template(document_path: str, replacements: Dict[str, str], output_path: str) -> dict:
    document = _load_document(document_path)
    replacement_count = 0

    for block in _iter_text_blocks(document):
        text = block.text
        updated = text
        for key, value in (replacements or {}).items():
            key_text = str(key)
            if "{{" in key_text:
                patterns = [key_text]
            else:
                patterns = [f"{{{{{key_text}}}}}", f"{{{{ {key_text} }}}}"]
            for pattern in patterns:
                if pattern in updated:
                    updated = updated.replace(pattern, str(value))
        if updated != text:
            block.text = updated
            replacement_count += 1

    save_path = output_path or _default_output_path(document_path, "filled")
    Path(save_path).parent.mkdir(parents=True, exist_ok=True)
    document.save(save_path)
    return {
        "operation": "fill_template",
        "document_name": Path(document_path).name,
        "replacement_count": replacement_count,
        "output_path": save_path,
    }


def _extract_outline(document_path: str) -> dict:
    document = _load_document(document_path)
    outline = []
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if not text:
            continue
        style_name = str(getattr(paragraph.style, "name", "") or "")
        if style_name.lower().startswith("heading"):
            match = re.search(r"(\d+)", style_name)
            level = int(match.group(1)) if match else 1
            outline.append({"text": text, "level": level, "style": style_name})
    return {
        "operation": "extract_outline",
        "document_name": Path(document_path).name,
        "items": outline,
        "count": len(outline),
    }


def _extract_tables(document_path: str) -> dict:
    document = _load_document(document_path)
    tables = []
    for index, table in enumerate(document.tables, start=1):
        rows = []
        for row in table.rows:
            rows.append([cell.text.strip() for cell in row.cells])
        tables.append({"table_index": index, "row_count": len(rows), "rows": rows})
    return {
        "operation": "extract_tables",
        "document_name": Path(document_path).name,
        "tables": tables,
        "table_count": len(tables),
    }


async def run(
    operation: str,
    document_path: str,
    keyword: str = "",
    replacement: str = "",
    output_path: str = "",
    replacements: Dict[str, str] | None = None,
) -> dict:
    if not operation:
        return error_result("operation is required")
    if operation not in SUPPORTED_OPERATIONS:
        return error_result(
            "Unsupported word operation",
            error=f"operation must be one of {sorted(SUPPORTED_OPERATIONS)}",
        )
    if not document_path:
        return error_result("document_path is required")

    try:
        if operation == "summarize":
            result = await asyncio.to_thread(_summarize_document, document_path)
            return success_result(data=result)

        if operation == "replace_text":
            if not keyword:
                return error_result("keyword is required when operation is replace_text")
            result = await asyncio.to_thread(_replace_text, document_path, keyword, replacement, output_path)
            return success_result(data=result)

        if operation == "fill_template":
            if not replacements:
                return error_result("replacements is required when operation is fill_template")
            result = await asyncio.to_thread(_fill_template, document_path, replacements, output_path)
            return success_result(data=result)

        if operation == "extract_outline":
            result = await asyncio.to_thread(_extract_outline, document_path)
            return success_result(data=result)

        if operation == "extract_tables":
            result = await asyncio.to_thread(_extract_tables, document_path)
            return success_result(data=result)

        result = await asyncio.to_thread(_format_document, document_path, output_path)
        return success_result(data=result)
    except FileNotFoundError as exc:
        return error_result("Word file does not exist", error=str(exc))
    except Exception as exc:
        return error_result("Failed to process Word document", error=str(exc))
